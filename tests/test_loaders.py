"""Loaders on synthetic files: PDF (incl. doubled text layer), DOCX, text, registry."""

from pathlib import Path

import docx
import pytest

from core.ingestion.cleaner import CleaningOptions, clean_pages, looks_character_doubled
from core.ingestion.loaders.base import LoaderRegistry
from core.ingestion.loaders.docx import DocxLoader
from core.ingestion.loaders.pdf import PdfLoader
from core.ingestion.loaders.registry import default_registry
from core.ingestion.loaders.text import TextLoader


def test_pdf_loader_pages_and_text(synthetic_pdf: Path):
    pages = PdfLoader().load(synthetic_pdf)
    assert [p.page for p in pages] == [1, 2, 3, 4, 5]
    assert "Acme Holdings plc Annual Report 2023" in pages[0].text
    assert "Revenue 1,234 million" in pages[0].text


def test_doubled_text_layer_is_repaired_by_cleaner(doubled_pdf: Path):
    pages = PdfLoader().load(doubled_pdf)
    assert "RReevveennuuee" in pages[0].text  # the defect as extracted
    cleaned = clean_pages(
        pages, CleaningOptions(min_chunk_chars=10, min_pages_for_header_detection=99)
    )
    assert "Revenue 9,876 million" in cleaned[0].text
    assert not looks_character_doubled(cleaned[0].text)


def test_duplicate_text_layer_runs_are_dropped(duplicate_layer_pdf: Path):
    naive = PdfLoader(drop_duplicate_layers=False).load(duplicate_layer_pdf)
    fixed = PdfLoader().load(duplicate_layer_pdf)
    assert len(naive[0].text) > 1.5 * len(fixed[0].text)
    assert fixed[0].text.count("Line 5 of the body") == 1
    assert "Layered Report page 1" in fixed[0].text
    # healthy pages are untouched
    healthy = PdfLoader().load(duplicate_layer_pdf.parent / "layered.pdf")
    assert all("LLiinnee" not in p.text for p in healthy)


def test_docx_loader_splits_on_page_breaks(tmp_path: Path):
    document = docx.Document()
    document.add_paragraph("First page text about markets.")
    document.add_paragraph("More on page one.")
    para = document.add_paragraph()
    para.add_run().add_break(docx.enum.text.WD_BREAK.PAGE)
    document.add_paragraph("Second page text.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "metric", "value"
    table.cell(1, 0).text, table.cell(1, 1).text = "revenue", "10"
    path = tmp_path / "news.docx"
    document.save(path)
    pages = DocxLoader().load(path)
    assert len(pages) == 2
    assert "First page" in pages[0].text and "Second page" in pages[1].text
    assert pages[0].tables == [[["metric", "value"], ["revenue", "10"]]]


def test_text_loader_form_feed_pages(tmp_path: Path):
    path = tmp_path / "a.txt"
    path.write_text("one\fthree", encoding="utf-8")
    pages = TextLoader().load(path)
    assert [(p.page, p.text) for p in pages] == [(1, "one"), (2, "three")]


def test_registry_lookup_and_unknown_extension():
    reg = default_registry()
    assert isinstance(reg.get("pdf"), PdfLoader)
    assert isinstance(reg.get(".DOCX"), DocxLoader)
    assert reg.supports(Path("x.md"))
    with pytest.raises(KeyError):
        reg.get(".xlsx")
    assert LoaderRegistry().extensions == []
