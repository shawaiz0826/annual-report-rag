"""DOCX loader built on python-docx.

Word documents have no fixed pagination; we split on explicit page breaks when present
and otherwise treat the whole document as page 1. Tables are returned as cell grids.
"""

from __future__ import annotations

from pathlib import Path

import docx
from docx.oxml.ns import qn

from core.models import PageText


class DocxLoader:
    """Load a .docx into one or more pages."""

    extensions = (".docx",)

    def load(self, path: Path) -> list[PageText]:
        """Read paragraphs (splitting on hard page breaks) and tables."""
        document = docx.Document(str(path))
        pages: list[list[str]] = [[]]
        for paragraph in document.paragraphs:
            if self._has_page_break(paragraph):
                pages.append([])
            text = paragraph.text.strip()
            if text:
                pages[-1].append(text)
        tables = [
            [[cell.text.strip() for cell in row.cells] for row in table.rows]
            for table in document.tables
        ]
        result = [
            PageText(page=i, text="\n".join(lines), tables=tables if i == 1 else [])
            for i, lines in enumerate(pages, start=1)
            if lines or i == 1
        ]
        return result or [PageText(page=1, text="", tables=tables)]

    @staticmethod
    def _has_page_break(paragraph: docx.text.paragraph.Paragraph) -> bool:
        for run in paragraph.runs:
            for br in run._element.findall(qn("w:br")):
                if br.get(qn("w:type")) == "page":
                    return True
        return False
